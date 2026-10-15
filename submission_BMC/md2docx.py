#!/usr/bin/env python3
"""Minimal Markdown -> DOCX converter tailored for manuscript_IBM.md."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = sys.argv[1] if len(sys.argv) > 1 else "manuscript_IBM.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else "manuscript_IBM.docx"

doc = Document()

# Base styles
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)

def add_inline(paragraph, text):
    """Apply **bold**, `code`, *italic* inline."""
    # Tokenize by the three markers, preserving order
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
    pos = 0
    for m in pattern.finditer(text):
        if pos < m.start():
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        run = paragraph.add_run(tok.strip("*").strip("`").strip("*"))
        if tok.startswith("**"):
            run.bold = True
        elif tok.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            run.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])

lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
n = len(lines)
in_code = False
code_buf = []
first_heading_done = False

while i < n:
    line = lines[i]
    # Code fence
    if line.strip().startswith("```"):
        if not in_code:
            in_code = True
            code_buf = []
        else:
            in_code = False
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run("\n".join(code_buf))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            code_buf = []
        i += 1
        continue
    if in_code:
        code_buf.append(line)
        i += 1
        continue

    s = line.strip()
    if s == "":
        i += 1
        continue
    if s == "---":
        i += 1
        continue

    # Table
    if s.startswith("|"):
        tbl_lines = []
        while i < n and lines[i].strip().startswith("|"):
            tbl_lines.append(lines[i].strip())
            i += 1
        # remove separator row(s)
        rows = [r for r in tbl_lines if not re.fullmatch(r"\|[\s:\-|]+\|", r)]
        if rows:
            cells = [r.strip("|").split("|") for r in rows]
            nc = max(len(c) for c in cells)
            table = doc.add_table(rows=len(cells), cols=nc)
            table.style = "Table Grid"
            for ri, row in enumerate(cells):
                for ci in range(nc):
                    val = row[ci].strip() if ci < len(row) else ""
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    r = cell.paragraphs[0].add_run(val)
                    if ri == 0:
                        r.bold = True
        continue

    # Title (# )
    if s.startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(s[2:].strip())
        run.bold = True
        run.font.size = Pt(15)
        i += 1
        continue

    # Heading 1 (## )
    if s.startswith("## "):
        p = doc.add_heading(level=1)
        add_inline(p, s[3:].strip())
        i += 1
        continue

    # Heading 2 (### )
    if s.startswith("### "):
        p = doc.add_heading(level=2)
        add_inline(p, s[4:].strip())
        i += 1
        continue

    # Bullet (- or *)
    if re.match(r"^[-*] ", s):
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, s[2:].strip())
        i += 1
        continue

    # Normal paragraph
    p = doc.add_paragraph()
    add_inline(p, s)
    i += 1

doc.save(OUT)
print("Wrote", OUT)
